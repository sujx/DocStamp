"""Office document property modification.

Supports .docx, .xlsx, .pptx formats.
Office docs are ZIP archives with metadata in docProps/core.xml (Dublin Core),
read directly when inspecting and written through the format's own library
(python-docx / openpyxl / python-pptx) when modifying.

Supports:
- Single file property modification
- Batch modification (multiple files, same properties)
- Unified time mode: set both created/modified to the same timestamp
- Separate time mode: set created and modified independently
"""

import os
import zipfile
from datetime import datetime
from xml.etree import ElementTree as ET

from errors import ErrorCode, ServiceResult


def read_properties(filepath: str) -> ServiceResult[dict]:
    """Read core properties from an Office document.

    Args:
        filepath: Path to the Office document (.docx/.xlsx/.pptx).

    Returns:
        ServiceResult with dict: created, modified, creator, last_modified_by.
    """
    props = {
        "created": "",
        "modified": "",
        "creator": "",
        "last_modified_by": "",
    }

    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            if "docProps/core.xml" not in zf.namelist():
                return ServiceResult.ok(props)

            xml_data = zf.read("docProps/core.xml")
            root = ET.fromstring(xml_data)

            # XML namespaces
            ns = {
                "dc": "http://purl.org/dc/elements/1.1/",
                "dcterms": "http://purl.org/dc/terms/",
                "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
            }

            # Dublin Core terms
            creator_el = root.find(".//dc:creator", ns)
            if creator_el is not None:
                props["creator"] = creator_el.text or ""

            modified_el = root.find(".//dcterms:modified", ns)
            if modified_el is not None:
                props["modified"] = modified_el.text or ""

            created_el = root.find(".//dcterms:created", ns)
            if created_el is not None:
                props["created"] = created_el.text or ""

            # OPC core properties
            last_mod_el = root.find(".//cp:lastModifiedBy", ns)
            if last_mod_el is not None:
                props["last_modified_by"] = last_mod_el.text or ""

    except (FileNotFoundError, zipfile.BadZipFile, ET.ParseError, KeyError) as e:
        return ServiceResult.fail(ErrorCode.UNSUPPORTED_FORMAT, f"Failed to read properties: {e}")

    return ServiceResult.ok(props)


def modify_properties(filepath: str, output_path: str, props: dict) -> ServiceResult[None]:
    """Modify core properties of an Office document.

    Args:
        filepath: Path to the source Office document.
        output_path: Path to write the modified document.
        props: dict with optional keys: created, modified, creator,
               last_modified_by.

    Returns:
        ServiceResult with None on success.
    """
    ext = os.path.splitext(filepath)[1].lower()

    try:
        if ext == ".docx":
            _modify_docx(filepath, output_path, props)
        elif ext == ".xlsx":
            _modify_xlsx(filepath, output_path, props)
        elif ext == ".pptx":
            _modify_pptx(filepath, output_path, props)
        else:
            return ServiceResult.fail(ErrorCode.UNSUPPORTED_FORMAT, f"Unsupported file format: {ext}")
    except Exception as e:
        return ServiceResult.fail(ErrorCode.VALIDATION_ERROR, str(e))

    return ServiceResult.ok(None)


def _modify_docx(filepath: str, output_path: str, props: dict) -> None:
    """Modify DOCX properties using python-docx API."""
    from docx import Document

    doc = Document(filepath)

    # Access core properties
    cp = doc.core_properties

    if "creator" in props:
        cp.author = str(props["creator"])
    if "last_modified_by" in props:
        cp.last_modified_by = str(props["last_modified_by"])
    if "created" in props:
        cp.created = _parse_datetime(props["created"])
    if "modified" in props:
        cp.modified = _parse_datetime(props["modified"])

    doc.save(output_path)


def _modify_xlsx(filepath: str, output_path: str, props: dict) -> None:
    """Modify XLSX properties using openpyxl API."""
    from openpyxl import load_workbook

    wb = load_workbook(filepath)

    if "creator" in props:
        wb.properties.creator = str(props["creator"])
    if "last_modified_by" in props:
        wb.properties.lastModifiedBy = str(props["last_modified_by"])
    if "created" in props:
        wb.properties.created = _parse_datetime(props["created"])
    if "modified" in props:
        wb.properties.modified = _parse_datetime(props["modified"])

    wb.save(output_path)


def _modify_pptx(filepath: str, output_path: str, props: dict) -> None:
    """Modify PPTX properties using python-pptx API."""
    from pptx import Presentation

    prs = Presentation(filepath)

    cp = prs.core_properties

    if "creator" in props:
        cp.author = str(props["creator"])
    if "last_modified_by" in props:
        cp.last_modified_by = str(props["last_modified_by"])
    if "created" in props:
        cp.created = _parse_datetime(props["created"])
    if "modified" in props:
        cp.modified = _parse_datetime(props["modified"])

    prs.save(output_path)


def _parse_datetime(value) -> datetime:
    """Parse a datetime string into a datetime object.

    Accepts ISO 8601 format or common date formats.
    """
    if isinstance(value, datetime):
        return value

    value = str(value).strip()

    # Try ISO 8601 formats
    for fmt in [
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%SZ",
        "%Y-%m-%dT%H:%M:%S%z",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d",
    ]:
        try:
            # Handle timezone suffix
            if value.endswith("Z"):
                value_clean = value[:-1] + "+00:00"
            else:
                value_clean = value
            return datetime.strptime(value_clean, fmt)
        except ValueError:
            continue

    raise ValueError(f"Cannot parse datetime: {value}")


def batch_modify_properties(
    filepaths: list,
    output_dir: str,
    props: dict,
    unify_time: bool = False,
) -> ServiceResult[list]:
    """Modify properties for multiple Office documents in batch.

    Args:
        filepaths: List of paths to Office documents.
        output_dir: Directory to write modified files.
        props: Properties dict.
        unify_time: If True, use the same timestamp for both created and modified.

    Returns:
        ServiceResult with list of dicts: [{"filename": "a.docx", "success": True}, ...]
    """
    if not filepaths:
        return ServiceResult.fail(ErrorCode.VALIDATION_ERROR, "No files provided for batch processing")

    results = []

    for filepath in filepaths:
        filename = os.path.basename(filepath)
        try:
            resolved_props = _resolve_time_props(props, unify_time)
            output_path = os.path.join(output_dir, filename)
            # modify_properties reports failures in its result rather than raising,
            # so an unchecked call would silently mark every file as a success.
            result = modify_properties(filepath, output_path, resolved_props)
            if result.success:
                results.append({"filename": filename, "success": True})
            else:
                results.append({
                    "filename": filename,
                    "success": False,
                    "error": result.message,
                })
        except Exception as e:
            results.append({"filename": filename, "success": False, "error": str(e)})

    return ServiceResult.ok(results)


def _resolve_time_props(props: dict, unify_time: bool) -> dict:
    """Resolve time properties based on unified/separate mode.

    Unified mode (unify_time=True):
        - If 'unified_time' is provided, use it for both created and modified.
        - If only 'created' or 'modified' is provided, copy to the other.
        - If neither is provided, no time change.

    Separate mode (unify_time=False):
        - created and modified are set independently.
        - Each is only set if explicitly provided.
    """
    result = dict(props)

    if unify_time:
        unified = props.get("unified_time") or props.get("created") or props.get("modified")
        if unified:
            result["created"] = unified
            result["modified"] = unified
            result.pop("unified_time", None)

    return result
