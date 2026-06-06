"""Excel/CSV file merging — combine multiple .xlsx/.csv files with the same structure."""

import csv
import os
from typing import List

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


class StructureMismatchError(ValueError):
    """Raised when Excel files have incompatible structures."""

    def __init__(self, filename: str, expected: List[str], actual: List[str]):
        self.filename = filename
        self.expected = expected
        self.actual = actual
        super().__init__(
            f"Structure mismatch in '{filename}': "
            f"expected columns {expected}, got {actual}"
        )


def _is_csv(filepath: str) -> bool:
    """Check if a file is a CSV based on extension."""
    return os.path.splitext(filepath)[1].lower() == ".csv"


def _get_headers(filepath: str) -> List[str]:
    """Extract column headers (first row values) from an Excel or CSV file."""
    if _is_csv(filepath):
        return _get_headers_csv(filepath)
    return _get_headers_xlsx(filepath)


def _get_headers_xlsx(filepath: str) -> List[str]:
    """Extract headers from an .xlsx file."""
    wb = load_workbook(filepath, read_only=True)
    try:
        ws = wb.active
        headers = []
        for cell in ws[1]:
            if cell.value is not None:
                headers.append(str(cell.value).strip())
        return headers
    finally:
        wb.close()


def _get_headers_csv(filepath: str) -> List[str]:
    """Extract headers from a CSV file."""
    with open(filepath, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        try:
            row = next(reader)
        except StopIteration:
            return []
        return [str(cell).strip() for cell in row if cell]


def merge_excel_files(
    filepaths: List[str],
    output_path: str,
) -> dict:
    """Merge multiple .xlsx/.csv files with the same structure into one .xlsx.

    All files must have the same column headers (based on first row).
    Rows from all files are appended in order, with a single header row.
    The output is always .xlsx format.

    Args:
        filepaths: List of paths to .xlsx or .csv files to merge.
        output_path: Path where the merged .xlsx will be written.

    Returns:
        dict with keys: total_rows, file_count, columns, per_file_rows.

    Raises:
        StructureMismatchError: If any file has different headers from the first.
        ValueError: If fewer than 2 files provided.
    """
    if len(filepaths) < 2:
        raise ValueError("At least 2 files are required for merging")

    # Validate all files exist
    for fp in filepaths:
        if not os.path.isfile(fp):
            raise ValueError(f"File not found: {fp}")

    # Get headers from the first file as the reference
    ref_headers = _get_headers(filepaths[0])
    if not ref_headers:
        raise ValueError(f"No headers found in '{os.path.basename(filepaths[0])}'")

    # Validate all other files have the same headers
    for fp in filepaths[1:]:
        headers = _get_headers(fp)
        if headers != ref_headers:
            raise StructureMismatchError(
                os.path.basename(fp), ref_headers, headers
            )

    # Create a new workbook and write the header row
    wb = Workbook()
    ws = wb.active
    ws.append(ref_headers)

    per_file_rows = []

    # Append data rows from each file
    for fp in filepaths:
        if _is_csv(fp):
            data_rows = _append_csv_rows(ws, fp)
        else:
            data_rows = _append_xlsx_rows(ws, fp)
        per_file_rows.append(data_rows)

    total_rows = sum(per_file_rows)
    wb.save(output_path)
    wb.close()

    return {
        "total_rows": total_rows,
        "file_count": len(filepaths),
        "columns": ref_headers,
        "per_file_rows": per_file_rows,
    }


def _append_xlsx_rows(ws, filepath: str) -> int:
    """Append data rows (skipping header) from an .xlsx file. Returns row count."""
    src_wb = load_workbook(filepath, read_only=True)
    try:
        src_ws = src_wb.active
        data_rows = 0
        for row in src_ws.iter_rows(min_row=2, values_only=True):
            if any(cell is not None for cell in row):
                ws.append(list(row))
                data_rows += 1
        return data_rows
    finally:
        src_wb.close()


def _append_csv_rows(ws, filepath: str) -> int:
    """Append data rows (skipping header) from a CSV file. Returns row count."""
    with open(filepath, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        try:
            next(reader)  # skip header
        except StopIteration:
            return 0
        data_rows = 0
        for row in reader:
            if any(cell for cell in row):
                ws.append(list(row))
                data_rows += 1
        return data_rows
